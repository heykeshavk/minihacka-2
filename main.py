from __future__ import annotations

import io
import re
import html
from collections import Counter
from datetime import datetime, timezone
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, field_validator

APP_VERSION = "1.0.0"
MAX_RESUME_BYTES = 8 * 1024 * 1024
MAX_RESUME_CHARS = 80_000
MAX_JD_CHARS = 25_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

app = FastAPI(
    title="ProofHire AI — Resume/JD Matching API",
    version=APP_VERSION,
    description=(
        "Evidence-first resume and job-description comparison. Scores are heuristic "
        "text-alignment indicators, not hiring probabilities or official ATS scores."
    ),
)

# For production, replace this list with your actual frontend origin(s).
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)

# Curated aliases only. This is intentionally deterministic and explainable.
SKILL_ALIASES: dict[str, list[str]] = {
    "python": ["python", "python3"],
    "java": ["java"],
    "javascript": ["javascript", "ecmascript"],
    "typescript": ["typescript"],
    "c++": ["c++", "cpp"],
    "c#": ["c#", "c sharp"],
    "sql": ["sql", "structured query language"],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "scikit-learn": ["scikit-learn", "sklearn"],
    "machine learning": ["machine learning", "ml"],
    "deep learning": ["deep learning"],
    "data analysis": ["data analysis", "data analytics"],
    "data visualization": ["data visualization", "data viz"],
    "statistics": ["statistics", "statistical analysis"],
    "excel": ["excel", "microsoft excel"],
    "power bi": ["power bi"],
    "tableau": ["tableau"],
    "fastapi": ["fastapi"],
    "flask": ["flask"],
    "django": ["django"],
    "rest api": ["rest api", "restful api", "restful services"],
    "api development": ["api development", "api integration"],
    "html": ["html", "html5"],
    "css": ["css", "css3"],
    "react": ["react", "react.js", "reactjs"],
    "next.js": ["next.js", "nextjs"],
    "node.js": ["node.js", "nodejs"],
    "git": ["git", "github", "gitlab"],
    "docker": ["docker", "containerization"],
    "aws": ["aws", "amazon web services"],
    "azure": ["azure", "microsoft azure"],
    "gcp": ["gcp", "google cloud platform"],
    "postgresql": ["postgresql", "postgres"],
    "mysql": ["mysql"],
    "mongodb": ["mongodb", "mongo db"],
    "data structures": ["data structures", "dsa"],
    "algorithms": ["algorithms", "algorithm design"],
    "problem solving": ["problem solving", "problem-solving"],
    "communication": ["communication", "verbal communication"],
    "teamwork": ["teamwork", "team collaboration", "collaboration"],
    "project management": ["project management"],
    "linux": ["linux"],
    "testing": ["testing", "unit testing", "pytest"],
    "nlp": ["nlp", "natural language processing"],
    "computer vision": ["computer vision", "opencv"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch", "torch"],
    "airflow": ["airflow", "apache airflow"],
    "spark": ["apache spark", "pyspark"],
    "powerpoint": ["powerpoint", "power point"],
    "communication skills": ["communication skills"],
    "agile": ["agile", "scrum"],
    "ci/cd": ["ci/cd", "continuous integration", "continuous delivery"],
    "kubernetes": ["kubernetes", "k8s"],
    "redis": ["redis"],
    "machine vision": ["machine vision"],
}

LEARNING_TIPS = {
    "python": "Build a small Python project and include a README with setup and usage instructions.",
    "sql": "Practice SELECT, JOIN, GROUP BY, subqueries, and window functions on a sample database.",
    "machine learning": "Train a baseline model, use a held-out test set, and report suitable evaluation metrics.",
    "data analysis": "Complete an exploratory analysis and explain three findings with evidence.",
    "data visualization": "Create a chart with clear labels, units, and a written interpretation.",
    "fastapi": "Build typed REST endpoints with validation, error handling, and automated tests.",
    "rest api": "Practice HTTP methods, status codes, request validation, and pagination.",
    "git": "Demonstrate branching, meaningful commits, pull requests, and conflict resolution.",
    "docker": "Containerize a small application and document its ports, environment variables, and volumes.",
    "aws": "Study IAM, compute, storage, networking, and deployment fundamentals.",
    "react": "Build a form-driven component using state, validation, and API integration.",
    "communication": "Prepare a concise project explanation covering context, actions, and measurable results.",
    "problem solving": "Solve progressively harder problems and explain the reasoning and complexity.",
}

STOP_WORDS = {
    "a", "an", "the", "and", "or", "but", "if", "then", "than", "to", "of",
    "in", "on", "for", "with", "from", "by", "as", "at", "is", "are", "be",
    "will", "would", "can", "could", "should", "must", "have", "has", "had",
    "we", "you", "your", "our", "their", "this", "that", "these", "those",
    "it", "its", "into", "about", "across", "within", "using", "use", "work",
    "working", "experience", "years", "year", "ability", "skills", "skill",
    "responsibilities", "responsibility", "requirements", "required", "preferred",
    "candidate", "role", "job", "description", "team", "strong", "good",
    "including", "etc", "such", "other", "new", "related", "support", "develop",
    "development", "build", "building", "maintain", "maintaining", "ensure",
    "provide", "manage", "management", "knowledge", "familiarity",
}

class AnalyzeTextRequest(BaseModel):
    resume_text: str = Field(min_length=30, max_length=MAX_RESUME_CHARS)
    job_description: str = Field(min_length=30, max_length=MAX_JD_CHARS)

    @field_validator("resume_text", "job_description")
    @classmethod
    def strip_required_text(cls, value: str) -> str:
        value = value.strip()
        if len(value) < 30:
            raise ValueError("Text must contain at least 30 non-edge characters.")
        return value


class ReportRequest(BaseModel):
    analysis: dict[str, Any]


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.casefold()).strip()


def _term_present(text: str, term: str) -> bool:
    """Boundary-aware matching, with explicit aliases for known tricky terms."""
    escaped = re.escape(term.casefold())
    return re.search(r"(?<![a-z0-9])" + escaped + r"(?![a-z0-9])", text.casefold()) is not None


def extract_skills(text: str) -> list[str]:
    normalized = _norm(text)
    found: list[str] = []
    for canonical, aliases in SKILL_ALIASES.items():
        if any(_term_present(normalized, alias) for alias in aliases):
            found.append(canonical)
    return sorted(found)


def split_sentences(text: str) -> list[str]:
    # Preserve evidence text as written; no generated claims are added.
    chunks = re.split(r"(?<=[.!?])\s+|\n+|(?<=;)\s+", text)
    return [re.sub(r"\s+", " ", x).strip(" \t•-*") for x in chunks if x.strip()]


def evidence_for_skill(resume_text: str, skill: str, limit: int = 3) -> list[dict[str, str]]:
    aliases = SKILL_ALIASES.get(skill, [skill])
    matches = []
    for sentence in split_sentences(resume_text):
        if any(_term_present(sentence, alias) for alias in aliases):
            matches.append({"quote": sentence[:600], "source": "resume_text"})
        if len(matches) >= limit:
            break
    return matches


def extract_keywords(text: str, limit: int = 20) -> list[dict[str, Any]]:
    tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9+#./-]{1,}", text.casefold())
    filtered = [
        t.strip("./-") for t in tokens
        if len(t.strip("./-")) >= 3 and t.strip("./-") not in STOP_WORDS
        and not t.strip("./-").isdigit()
    ]
    counts = Counter(filtered)
    return [{"keyword": word, "count": count} for word, count in counts.most_common(limit)]


def cosine_similarity_score(resume_text: str, jd_text: str) -> tuple[float | None, str]:
    """
    TF-IDF cosine similarity if scikit-learn is installed; deterministic token overlap
    fallback otherwise. The method is always disclosed in the response.
    """
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity

        matrix = TfidfVectorizer(ngram_range=(1, 2), stop_words="english").fit_transform(
            [resume_text, jd_text]
        )
        value = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])
        return round(value * 100, 2), "tfidf_cosine"
    except ImportError:
        r = set(re.findall(r"\b[a-zA-Z][a-zA-Z0-9+#.-]{2,}\b", resume_text.casefold()))
        j = set(re.findall(r"\b[a-zA-Z][a-zA-Z0-9+#.-]{2,}\b", jd_text.casefold()))
        if not r or not j:
            return None, "unavailable"
        return round(100 * len(r & j) / len(r | j), 2), "token_jaccard_fallback"
    except ValueError:
        # Empty vocabulary or all stopwords; no misleading fabricated score.
        return None, "unavailable"


def ats_format_checks(resume_text: str, filename: str | None) -> list[dict[str, Any]]:
    checks = []
    text = resume_text.strip()
    checks.append({
        "id": "text_extracted",
        "label": "Readable text extraction",
        "status": "pass" if len(text) >= 100 else "warning",
        "detail": "Text was extracted from the provided file." if len(text) >= 100
                  else "Very little text was extracted; verify the file or use a text-based PDF/DOCX.",
    })
    has_email = bool(re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I))
    checks.append({
        "id": "contact_email",
        "label": "Email address detected",
        "status": "pass" if has_email else "warning",
        "detail": "An email-like string was detected." if has_email
                  else "No email address detected in extracted text; add one if appropriate.",
    })
    has_phone = bool(re.search(r"(?<!\w)(?:\+?\d[\d ()-]{7,}\d)(?!\w)", text))
    checks.append({
        "id": "contact_phone",
        "label": "Phone number detected",
        "status": "pass" if has_phone else "warning",
        "detail": "A phone-like number was detected." if has_phone
                  else "No phone-like number detected; add a contact number if appropriate.",
    })
    headings = ["education", "experience", "work experience", "projects", "skills", "certifications"]
    found_headings = [h for h in headings if re.search(r"(?im)^\s*" + re.escape(h) + r"\s*:?\s*$", text)]
    checks.append({
        "id": "section_headings",
        "label": "Common section headings",
        "status": "pass" if len(found_headings) >= 2 else "warning",
        "detail": "Detected headings: " + (", ".join(found_headings) if found_headings else "none detected"),
    })
    checks.append({
        "id": "file_type",
        "label": "File type",
        "status": "pass" if filename and filename.lower().endswith((".pdf", ".docx", ".txt")) else "warning",
        "detail": "Supported file type." if filename and filename.lower().endswith((".pdf", ".docx", ".txt"))
                  else "File extension could not be confirmed.",
    })
    return checks


def ats_friendliness_score(checks: list[dict[str, Any]]) -> dict[str, Any]:
    # This is a transparent heuristic checklist, not an actual ATS vendor score.
    passed = sum(c["status"] == "pass" for c in checks)
    score = round(100 * passed / len(checks)) if checks else 0
    return {
        "score": score,
        "label": "format_check_heuristic",
        "passed_checks": passed,
        "total_checks": len(checks),
        "disclaimer": "A basic format/readability checklist only; not a score from any ATS vendor.",
    }


def build_analysis(resume_text: str, job_description: str, filename: str | None = None) -> dict[str, Any]:
    resume_text = resume_text.strip()
    job_description = job_description.strip()
    if len(resume_text) < 30:
        raise HTTPException(status_code=422, detail="Resume text is too short or empty.")
    if len(job_description) < 30:
        raise HTTPException(status_code=422, detail="Job description must contain at least 30 characters.")

    resume_skills = set(extract_skills(resume_text))
    jd_skills = set(extract_skills(job_description))
    matched = sorted(resume_skills & jd_skills)
    missing = sorted(jd_skills - resume_skills)

    # A partial match is not inferred: the deterministic MVP reports only matched/missing.
    skill_details = []
    for skill in sorted(jd_skills):
        evidence = evidence_for_skill(resume_text, skill)
        skill_details.append({
            "skill": skill,
            "status": "matched" if skill in matched else "missing_keyword_evidence",
            "resume_evidence": evidence,
            "evidence_found": bool(evidence),
            "interpretation": (
                "A matching term was found in the resume. This does not verify practical ability."
                if evidence else
                "No matching term was found by the current skill dictionary. This does not prove the candidate lacks the skill."
            ),
        })

    jd_skill_count = len(jd_skills)
    skill_match_score = round(100 * len(matched) / jd_skill_count) if jd_skill_count else None
    cosine_score, cosine_method = cosine_similarity_score(resume_text, job_description)
    checks = ats_format_checks(resume_text, filename)
    ats_score = ats_friendliness_score(checks)

    roadmap = []
    for skill in missing[:10]:
        roadmap.append({
            "skill": skill,
            "tip": LEARNING_TIPS.get(skill, f"Build a small portfolio project that demonstrates {skill}, and document your approach."),
            "suggested_evidence": f"Add a truthful project, work example, or coursework entry demonstrating {skill} if you have one.",
        })

    warnings = []
    if not jd_skills:
        warnings.append(
            "No skills from the curated dictionary were detected in the job description. "
            "The skill-match score is unavailable; review the extracted keywords and extend the dictionary if needed."
        )
    if cosine_score is None:
        warnings.append("Text similarity could not be calculated because there was not enough usable vocabulary.")
    warnings.append("No external facts, job listings, or candidate claims were generated by this analysis.")

    return {
        "ok": True,
        "product": "ProofHire AI",
        "version": APP_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "resume_filename": filename,
        "summary": {
            "skill_match_score": skill_match_score,
            "skill_match_score_type": "matched_detected_skills / detected_job_skills",
            "matched_skills_count": len(matched),
            "missing_keyword_count": len(missing),
            "detected_job_skills_count": jd_skill_count,
            "detected_resume_skills_count": len(resume_skills),
            "cosine_similarity_score": cosine_score,
            "cosine_similarity_method": cosine_method,
            "ats_friendliness": ats_score,
            "disclaimer": (
                "These are transparent text-analysis heuristics, not hiring probabilities, "
                "proof of competence, or an official ATS score."
            ),
        },
        "skills": {
            "matched": matched,
            "missing_keyword_evidence": missing,
            "details": skill_details,
            "resume_skills_detected": sorted(resume_skills),
            "job_skills_detected": sorted(jd_skills),
        },
        "keyword_analysis": {
            "job_description_keywords": extract_keywords(job_description, 25),
            "resume_keywords": extract_keywords(resume_text, 25),
        },
        "ats_checks": checks,
        "improvement_tips": [
            {
                "id": "evidence",
                "title": "Add evidence, not just a skill list",
                "detail": "Where truthful, support relevant skills with project bullets, outcomes, coursework, or work examples.",
            },
            {
                "id": "tailor",
                "title": "Tailor the resume to this role",
                "detail": "Use relevant job-description terminology only where it accurately describes your experience.",
            },
            {
                "id": "readability",
                "title": "Keep formatting readable",
                "detail": "Use clear section headings and check that the extracted text preserves your intended content.",
            },
            {
                "id": "impact",
                "title": "Make project bullets specific",
                "detail": "Describe the action, tools used, and verifiable outcome. Do not invent metrics.",
            },
        ],
        "learning_roadmap": roadmap,
        "limitations": [
            "Skill detection uses a curated alias dictionary and may miss synonyms or domain-specific terms.",
            "Missing keyword evidence is not proof that a person lacks a skill.",
            "Text similarity measures textual overlap, not candidate quality or hiring likelihood.",
            "The ATS-friendliness value is a small transparent format checklist, not a vendor score.",
        ],
        "warnings": warnings,
        "analysis_method": {
            "skill_matching": "deterministic curated dictionary and boundary-aware matching",
            "text_similarity": cosine_method,
            "generative_ai_used": False,
        },
    }


async def extract_resume_text(filename: str, content: bytes) -> str:
    if not filename:
        raise HTTPException(status_code=400, detail="Resume filename is required.")
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded resume is empty.")
    if len(content) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume exceeds the 8 MB size limit.")

    try:
        if ext == ".txt":
            text = content.decode("utf-8-sig", errors="replace")
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content), strict=False)
            if len(reader.pages) > 40:
                raise HTTPException(status_code=413, detail="PDF exceeds the 40-page limit.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document
            document = Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Could not parse this file. Verify that it is a valid, unencrypted document.") from exc

    text = text.strip()
    if len(text) < 30:
        raise HTTPException(
            status_code=422,
            detail="Very little selectable text was extracted. If this is a scanned PDF, OCR is required; this backend does not pretend OCR succeeded.",
        )
    return text[:MAX_RESUME_CHARS]


@app.get("/")
def root():
    return {
        "name": "ProofHire AI Resume/JD Matching API",
        "version": APP_VERSION,
        "status": "running",
        "docs": "/docs",
        "health": "/health",
    }


@app.get("/health")
def health():
    return {"status": "ok", "version": APP_VERSION}


@app.get("/api/skills")
def skill_catalog():
    return {"skills": sorted(SKILL_ALIASES), "count": len(SKILL_ALIASES)}


@app.post("/api/analyze")
async def analyze_upload(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
):
    if len(job_description.strip()) < 30:
        raise HTTPException(status_code=422, detail="Job description must contain at least 30 characters.")
    content = await resume.read()
    text = await extract_resume_text(resume.filename or "", content)
    return build_analysis(text, job_description, resume.filename)


@app.post("/api/analyze-text")
def analyze_text(payload: AnalyzeTextRequest):
    return build_analysis(payload.resume_text, payload.job_description, "inline-text")


@app.post("/api/report")
def create_html_report(payload: ReportRequest):
    """Create a downloadable HTML report from a previous /api/analyze response."""
    data = payload.analysis
    summary = data.get("summary") or {}
    skills = data.get("skills") or {}
    matched = skills.get("matched") or []
    missing = skills.get("missing_keyword_evidence") or []
    checks = data.get("ats_checks") or []

    def esc(value: Any) -> str:
        return html.escape(str(value), quote=True)

    matched_items = "".join(f"<li>{esc(s)}</li>" for s in matched) or "<li>None detected</li>"
    missing_items = "".join(f"<li>{esc(s)}</li>" for s in missing) or "<li>None detected</li>"
    check_items = "".join(
        f"<li><strong>{esc(c.get('label', 'Check'))}</strong>: {esc(c.get('status', 'unknown'))} — {esc(c.get('detail', ''))}</li>"
        for c in checks
    )
    body = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>ProofHire AI Report</title>
<style>body{{font-family:Arial,sans-serif;max-width:900px;margin:32px auto;padding:0 20px;color:#17212b;line-height:1.55}}
h1{{color:#087f5b}}.score{{font-size:36px;font-weight:bold}}section{{border:1px solid #d8e1e7;border-radius:12px;padding:16px;margin:14px 0}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}small,.muted{{color:#52616b}}@media(max-width:600px){{.grid{{grid-template-columns:1fr}}}}</style></head><body>
<h1>ProofHire AI</h1><p class="muted">Resume and job-description comparison report</p>
<section><div class="score">{esc(summary.get('skill_match_score', 'N/A'))}{'%' if summary.get('skill_match_score') is not None else ''}</div>
<p>Detected skill overlap · cosine similarity: {esc(summary.get('cosine_similarity_score', 'N/A'))}{'%' if summary.get('cosine_similarity_score') is not None else ''}</p>
<p>{esc(summary.get('disclaimer', 'Heuristic results only.'))}</p></section>
<div class="grid"><section><h2>Matched skills</h2><ul>{matched_items}</ul></section>
<section><h2>Missing keyword evidence</h2><ul>{missing_items}</ul></section></div>
<section><h2>ATS readability checks</h2><ul>{check_items or '<li>No checks provided</li>'}</ul></section>
<section><h2>Interpretation</h2><p>Missing keyword evidence does not prove the candidate lacks a skill. Similarity and format checks are not hiring predictions or official ATS vendor scores.</p></section>
<small>Generated at {esc(data.get('generated_at', 'unknown time'))}. Report is based only on the submitted resume and job description.</small>
</body></html>"""
    return StreamingResponse(
        io.BytesIO(body.encode("utf-8")),
        media_type="text/html; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="proofhire-report.html"'},
    )


@app.get("/api/demo/sample")
def demo_sample():
    """Clearly labelled synthetic data for a quick demo; not scraped job-market data."""
    return {
        "label": "SYNTHETIC_DEMO_DATA_NOT_REAL_JOB_LISTINGS",
        "resume_text": (
            "Aarav Sharma\nEmail: aarav@example.com\nPhone: +91 98765 43210\n"
            "Skills\nPython, SQL, Pandas, FastAPI, Git\n"
            "Projects\nBuilt a Python data analysis project using Pandas and SQL. "
            "Created REST API endpoints with FastAPI and wrote basic tests.\n"
            "Education\nBSc Computer Science"
        ),
        "job_description": (
            "Junior Data Analyst / Python Developer. Required skills: Python, SQL, Pandas, "
            "data analysis, data visualization, FastAPI, Docker, Git. Responsibilities include "
            "building data pipelines, creating reports, testing APIs, and communicating findings."
        ),
        "note": "Synthetic example for testing only. It is not sourced from Indeed or a real employer.",
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
